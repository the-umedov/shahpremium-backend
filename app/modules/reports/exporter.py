"""exporter.js ekvivalenti — CSV/Excel/PDF eksport.

Original ExcelJS/PDFKit ishlatgan; bu portda openpyxl (xlsx) va reportlab (pdf)
ishlatiladi (requirements.txt'da mavjud). Ustunlar/qatorlar tarkibi bir xil,
faqat kutubxona-darajasidagi API farq qiladi.
"""

from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

from app.modules.reports.schemas import ReportData


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def to_csv(data: ReportData) -> bytes:
    def esc(s: str) -> str:
        return '"' + s.replace('"', '""') + '"'

    head = ",".join(esc(c.header) for c in data.columns)
    body = "\n".join(
        ",".join(esc(_cell(row.get(c.key))) for c in data.columns) for row in data.rows
    )
    # BOM — Excel'da UTF-8 to'g'ri ko'rinishi uchun
    return ("﻿" + head + "\n" + body).encode("utf-8")


def to_xlsx(data: ReportData) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (data.title[:30] or "Report")
    ws.append([c.header for c in data.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in data.rows:
        ws.append([_cell(row.get(c.key)) if not isinstance(row.get(c.key), (int, float)) else row.get(c.key) for c in data.columns])
    for i in range(1, len(data.columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(data: ReportData) -> bytes:
    buf = BytesIO()
    page_size = landscape(A4)
    width, height = page_size
    margin = 36
    c = canvas.Canvas(buf, pagesize=page_size)

    def new_page() -> float:
        c.setFont("Helvetica-Bold", 16)
        c.drawString(margin, height - margin, data.title)
        return height - margin - 24

    y = new_page()
    col_width = (width - 2 * margin) / max(1, len(data.columns))

    def draw_row(values: list[str], y_pos: float, bold: bool = False) -> None:
        c.setFont("Helvetica-Bold" if bold else "Helvetica", 9)
        # ellipsis emulyatsiyasi — uzun matnni ustun kengligiga moslab qirqamiz
        max_chars = max(4, int(col_width / 5))
        for i, v in enumerate(values):
            text = v if len(v) <= max_chars else v[: max_chars - 1] + "…"
            c.drawString(margin + i * col_width, y_pos, text)

    draw_row([col.header for col in data.columns], y, bold=True)
    y -= 4
    c.line(margin, y, width - margin, y)
    y -= 14

    for row in data.rows:
        if y < margin:
            c.showPage()
            y = new_page()
            c.line(margin, y - 4, width - margin, y - 4)
            y -= 18
        draw_row([_cell(row.get(col.key)) for col in data.columns], y)
        y -= 14

    c.save()
    return buf.getvalue()


def export_report(data: ReportData, fmt: str) -> tuple[bytes, str, str]:
    safe = re.sub(r"[^\w-]+", "_", data.title).lower() or "report"
    if fmt == "csv":
        return to_csv(data), "text/csv; charset=utf-8", f"{safe}.csv"
    if fmt == "xlsx":
        return (
            to_xlsx(data),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{safe}.xlsx",
        )
    return to_pdf(data), "application/pdf", f"{safe}.pdf"
